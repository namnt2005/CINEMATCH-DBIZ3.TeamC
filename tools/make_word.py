# -*- coding: utf-8 -*-
"""Convert docs/spec/*.md and docs/screens/*.md to Word (docs/word/specs, docs/word/screens) with a consistent style."""
import toolpaths
import glob, os, subprocess, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_ref.docx")
NAVY, TEAL, GREY = RGBColor(0x1F, 0x2A, 0x5C), RGBColor(0x1F, 0x5F, 0x8B), RGBColor(0x5B, 0x62, 0x6C)


def build_reference():
    subprocess.run([toolpaths.pandoc(), "-o", TMP, "--print-default-data-file", "reference.docx"], check=True)
    d = Document(TMP)
    st = d.styles
    def font(name, size, color=None, bold=None, italic=None, face="Calibri"):
        cand = [x for x in st if x.name and x.name.lower() == name.lower()]
        if not cand: return
        s = cand[0]; s.font.name = face; s.font.size = Pt(size)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), face) if s.element.rPr is not None and s.element.rPr.rFonts is not None else None
        if color is not None: s.font.color.rgb = color
        if bold is not None: s.font.bold = bold
        if italic is not None: s.font.italic = italic
    font("Normal", 10)
    for n in ("Body Text", "First Paragraph", "Compact"):
        if n in [x.name for x in st]: font(n, 10)
    font("Title", 20, NAVY, True)
    font("Heading 1", 15, NAVY, True)
    font("Heading 2", 12.5, NAVY, True)
    font("Heading 3", 11, TEAL, True)
    if "Block Text" in [x.name for x in st]: font("Block Text", 9, GREY, italic=True)
    font("Source Code", 8, face="Consolas")
    font("Verbatim Char", 8, face="Consolas")
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.8)
    d.save(TMP)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "C9CED6")
        b.append(e)
    tblPr.append(b)


def post(path):
    d = Document(path)
    for t in d.tables:
        borders(t)
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn("w:tblW")):
            tblPr.remove(old)
        w = OxmlElement("w:tblW"); w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct"); tblPr.append(w)
        grid = t._tbl.tblGrid
        cols = grid.findall(qn("w:gridCol")) if grid is not None else []
        tot = sum(int(c.get(qn("w:w")) or 0) for c in cols)
        if tot:
            f = 9860 / tot
            for c in cols:
                c.set(qn("w:w"), str(int(int(c.get(qn("w:w"))) * f)))
            for tc in t._tbl.iter(qn("w:tc")):
                tcw = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
                if tcw is not None and tcw.get(qn("w:type")) == "dxa":
                    tcw.set(qn("w:w"), str(int(int(tcw.get(qn("w:w"))) * f)))
        for i, row in enumerate(t.rows):
            for c in row.cells:
                if i == 0: shade(c, "E4ECF7")
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        if i == 0: r.font.bold = True; r.font.color.rgb = NAVY
    d.save(path)


def convert(md, out, resource):
    os.makedirs(os.path.dirname(out), exist_ok=True)
    r = subprocess.run([toolpaths.pandoc(), md, "-f", "gfm", "-o", out, "--reference-doc", TMP, "--resource-path", resource], capture_output=True, text=True)
    if r.returncode:
        print("FAIL", md, r.stderr[:300]); return
    post(out)


def data_word():
    """Word copies of the Session 5 data package into docs/word/data/ (03-erd.mmd is a diagram: see erd-views/)."""
    import re, tempfile
    out = f"{ROOT}/docs/word/data"; os.makedirs(out, exist_ok=True)
    t = tempfile.mkdtemp()
    names = {"SYS": "Platform foundation", "M1": "Segment router", "M0": "Project workspace", "M2": "Content and dossier checks",
             "M3": "Location discovery", "M4": "Vietnamese service partners", "M5": "Dossier kit", "M7": "VFDA support"}
    for name in ["01-entity-dictionary", "02-crud-matrix", "04-data-model", "05-review"]:
        md = re.sub(r"^---\n.*?\n---\n", "", open(f"{ROOT}/data/{name}.md", encoding="utf-8").read(), flags=re.S)
        if name == "04-data-model":
            views = "\n\n".join(f"**{m} — {n}** (relationships ending in a {m} entity)\n\n![{m}](erd-views/erd-{m}.png)" for m, n in names.items())
            md = re.sub(r"```mermaid\n.*?```", "The full diagram is too wide to print legibly (criterion 4 in 05-review.md). "
                        "It is shown here as one view per module; the lines are the same as in `data/03-erd.mmd`.\n\n" + views, md, flags=re.S)
        open(f"{t}/{name}.md", "w", encoding="utf-8").write(md)
        convert(f"{t}/{name}.md", f"{out}/{name}.docx", out)
    seed = re.sub(r"^---\n.*?\n---\n", "", open(f"{ROOT}/data/seed/README.md", encoding="utf-8").read(), flags=re.S)
    open(f"{t}/seed.md", "w", encoding="utf-8").write(seed)
    convert(f"{t}/seed.md", f"{out}/seed-README.docx", out)


if __name__ == "__main__":
    build_reference()
    data_word()
    for md in sorted(glob.glob(f"{ROOT}/docs/spec/spec-*.md")):
        convert(md, f"{ROOT}/docs/word/specs/" + os.path.basename(md)[:-3] + ".docx", f"{ROOT}/docs/spec")
    for md in sorted(glob.glob(f"{ROOT}/docs/screens/screen-spec-*.md")):
        convert(md, f"{ROOT}/docs/word/screens/" + os.path.basename(md)[:-3] + ".docx", f"{ROOT}/docs/screens")
    os.remove(TMP)
    print("done:", len(glob.glob(f"{ROOT}/docs/word/specs/*.docx")), "specs,", len(glob.glob(f"{ROOT}/docs/word/screens/*.docx")), "screens")
