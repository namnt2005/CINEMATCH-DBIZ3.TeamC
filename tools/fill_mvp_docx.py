# -*- coding: utf-8 -*-
"""Fill the Session 1 MVP Scope template (.docx) with docs/mvp-scope data (mvp.py)."""
import copy, sys
from docx import Document
import mvp

src, out = sys.argv[1], sys.argv[2]
d = Document(src)
T = d.tables


def set_cell(cell, text, bold=None):
    ps = cell.paragraphs
    for p in ps[1:]:
        p._element.getparent().remove(p._element)
    p = ps[0]
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        r = p.add_run(text)
        if bold is not None:
            r.bold = bold


def set_cell_lines(cell, lines):
    set_cell(cell, lines[0])
    base = cell.paragraphs[0]
    for ln in lines[1:]:
        np_ = copy.deepcopy(base._element)
        cell._element.append(np_)
        from docx.text.paragraph import Paragraph
        par = Paragraph(np_, cell)
        for r in par.runs[1:]:
            r._element.getparent().remove(r._element)
        par.runs[0].text = ln


# title label
p0 = d.paragraphs[0]
if p0.runs:
    p0.runs[0].text = "SESSION 1 DELIVERABLE · GROUP C · CINEMATCH"
    for r in p0.runs[1:]:
        r._element.getparent().remove(r._element)

# header
for c, v in zip(T[0].rows[1].cells, mvp.HEADER):
    set_cell(c, v)
# problem / target
set_cell(T[1].rows[0].cells[1], mvp.PROBLEM)
set_cell(T[2].rows[0].cells[1], mvp.TARGET)
# link to DBIZ2
for row, (_, lines) in zip(T[3].rows, mvp.LINK):
    set_cell_lines(row.cells[1], ["• " + x.replace("*", "").replace("`", "") for x in lines])

# MoSCoW: clone rows of the same priority so the coloured priority cell is kept
tbl = T[4]
proto = {}
for r in tbl.rows[1:]:
    proto.setdefault(r.cells[0].text.strip(), r)
for r in list(tbl.rows[1:]):
    r._tr.getparent().remove(r._tr)
for pr, feat, note in mvp.MOSCOW:
    tr = copy.deepcopy(proto[pr]._tr)
    tbl._tbl.append(tr)
    row = tbl.rows[-1]
    set_cell(row.cells[1], feat)
    set_cell(row.cells[2], note)

# backlog
bt = T[5]
proto_b = bt.rows[1]._tr
for r in list(bt.rows[1:]):
    r._tr.getparent().remove(r._tr)
for i, (item, pr, owner, status) in enumerate(mvp.BACKLOG, 1):
    bt._tbl.append(copy.deepcopy(proto_b))
    row = bt.rows[-1]
    for c, v in zip(row.cells, (str(i), item, pr, owner, status)):
        set_cell(c, v)
d.save(out)
print("saved", out)
